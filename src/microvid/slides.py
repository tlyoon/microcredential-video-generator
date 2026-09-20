from __future__ import annotations

import math
import re
from pathlib import Path
from zipfile import ZipFile

import yaml
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from .office_math import (
    DRAWING_2010_NS,
    DRAWING_NS,
    looks_like_math,
    set_native_math_paragraph,
    visible_math_to_latex,
)
from .qa import validate_slide_surface

_TEXT = RGBColor(28, 32, 36)
_MUTED = RGBColor(90, 98, 105)
_PANEL_FILL = RGBColor(245, 247, 249)
_PANEL_LINE = RGBColor(205, 211, 217)
_ACCENT = RGBColor(45, 94, 145)
_SLIDE_BOTTOM = 7.12
_STANDARD_CONTENT_TOP = 1.55
_TEXTBOOK_CONTENT_TOP = 0.65
_REGION_GAP = 0.18


def _configure_text_frame(
    text_frame,
    *,
    margin_x: float = 0.06,
    margin_y: float = 0.04,
    vertical_anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    """Keep rendered text inside its allocated region."""
    text_frame.word_wrap = True
    text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    text_frame.margin_left = Inches(margin_x)
    text_frame.margin_right = Inches(margin_x)
    text_frame.margin_top = Inches(margin_y)
    text_frame.margin_bottom = Inches(margin_y)
    text_frame.vertical_anchor = vertical_anchor


def _set_visible_paragraph(
    paragraph,
    text: str,
    *,
    font_size: int,
    color: RGBColor = _TEXT,
    bold: bool = False,
    alignment: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    """Render every visible mathematical expression as editable Office Math."""
    value = str(text).strip()
    if looks_like_math(value):
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        paragraph.alignment = alignment
        set_native_math_paragraph(
            paragraph,
            visible_math_to_latex(value),
            font_size_pt=font_size,
            alignment="centerGroup" if alignment == PP_ALIGN.CENTER else "left",
        )
        return
    paragraph.text = value
    paragraph.font.size = Pt(font_size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    paragraph.alignment = alignment
    for run in paragraph.runs:
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color


def _estimated_line_count(text: str, width_inches: float, font_size: int) -> int:
    usable_points = max(36.0, width_inches * 72.0 - 12.0)
    chars_per_line = max(10, int(usable_points / max(1.0, font_size * 0.56)))
    return max(1, math.ceil(len(str(text).strip()) / chars_per_line))


def _estimated_text_height(
    lines: list[str],
    width_inches: float,
    font_size: int,
    *,
    paragraph_gap_points: float = 5.0,
) -> float:
    if not lines:
        return 0.0
    points = 8.0
    for line in lines:
        line_factor = 1.35 if looks_like_math(str(line)) else 1.18
        points += _estimated_line_count(str(line), width_inches, font_size) * font_size * line_factor
        points += paragraph_gap_points
    return points / 72.0


def _lead_region_height(
    lines: list[str],
    width_inches: float,
    *,
    preferred_font_size: int,
    maximum: float,
) -> float:
    if not lines:
        return 0.0
    required = _estimated_text_height(lines, width_inches, preferred_font_size)
    return min(maximum, max(0.72, required))


def _title_font_size(title: str) -> int:
    n = len(title.strip())
    if n <= 42:
        return 30
    if n <= 68:
        return 26
    if n <= 95:
        return 23
    return 20


def _add_title(slide, title: str) -> None:
    box = slide.shapes.add_textbox(Inches(0.78), Inches(0.42), Inches(11.75), Inches(1.05))
    tf = box.text_frame
    tf.clear()
    _configure_text_frame(tf, margin_x=0.0, margin_y=0.0)
    p = tf.paragraphs[0]
    _set_visible_paragraph(
        p,
        str(title),
        font_size=_title_font_size(str(title)),
        bold=True,
    )


def _body_font_size(lines: list[str], *, has_figure: bool) -> int:
    chars = len(" ".join(str(x) for x in lines))
    if chars <= 90:
        return 28 if not has_figure else 25
    if chars <= 190:
        return 25 if not has_figure else 22
    if chars <= 300:
        return 22 if not has_figure else 19
    return 18


def _add_body(
    slide,
    item: dict,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    has_figure: bool,
    show_equation: bool = True,
) -> None:
    lines = [str(line) for line in item.get("onscreen", [])]
    if not lines and not (show_equation and item.get("equation_latex")):
        return
    body = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    btf = body.text_frame
    btf.clear()
    _configure_text_frame(btf)
    font_size = _body_font_size(lines, has_figure=has_figure)
    for j, line in enumerate(lines):
        p = btf.paragraphs[0] if j == 0 else btf.add_paragraph()
        _set_visible_paragraph(p, line, font_size=font_size)
        p.space_after = Pt(5)

    equation = item.get("equation_latex") if show_equation else None
    if equation:
        p = btf.add_paragraph()
        p.space_before = Pt(8)
        set_native_math_paragraph(
            p, str(equation), font_size_pt=max(22, font_size), alignment="centerGroup"
        )


def _add_equation_focus(
    slide,
    item: dict,
    *,
    content_top: float = _STANDARD_CONTENT_TOP,
) -> None:
    lines = [str(x) for x in item.get("onscreen", [])]
    body_y = content_top
    body_h = _lead_region_height(lines, 11.25, preferred_font_size=24, maximum=1.45)
    if lines:
        _add_body(
            slide,
            item,
            1.0,
            body_y,
            11.25,
            body_h,
            has_figure=False,
            show_equation=False,
        )
    eq = item.get("equation_latex")
    equation_y = body_y + body_h + (_REGION_GAP if body_h else 0.25)
    if eq:
        box = slide.shapes.add_textbox(Inches(1.05), Inches(equation_y), Inches(11.2), Inches(1.35))
        _configure_text_frame(
            box.text_frame,
            margin_x=0.02,
            margin_y=0.02,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        p = box.text_frame.paragraphs[0]
        set_native_math_paragraph(p, str(eq), font_size_pt=30, alignment="centerGroup")
    panels = item.get("visual_panels", []) or []
    if panels:
        panel_y = equation_y + (1.35 if eq else 0.0) + _REGION_GAP
        _add_panel_grid(
            slide,
            panels,
            1.1,
            panel_y,
            11.1,
            max(1.25, _SLIDE_BOTTOM - panel_y),
            process=False,
        )


def _add_equation_footer(slide, latex: str, *, top: float, height: float = 0.95) -> None:
    box = slide.shapes.add_textbox(Inches(1.05), Inches(top), Inches(11.2), Inches(height))
    _configure_text_frame(
        box.text_frame,
        margin_x=0.02,
        margin_y=0.0,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    paragraph = box.text_frame.paragraphs[0]
    set_native_math_paragraph(paragraph, str(latex), font_size_pt=18, alignment="centerGroup")


def _panel_text_size(panel_count: int, body_chars: int) -> int:
    if panel_count <= 2 and body_chars < 110:
        return 20
    if body_chars < 170:
        return 17
    return 15


def _add_panel_grid(
    slide,
    panels: list[dict],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    process: bool,
) -> None:
    if not panels:
        return
    count = min(len(panels), 6)
    panels = panels[:count]
    if (process and count <= 4) or count <= 2:
        cols, rows = count, 1
    elif count <= 4:
        cols, rows = 2, 2
    else:
        cols, rows = 3, 2
    gap = 0.18
    arrow_gap = 0.32 if process and rows == 1 and cols > 1 else 0
    usable_w = w - gap * (cols - 1) - arrow_gap * (cols - 1)
    cell_w = usable_w / cols
    cell_h = (h - gap * (rows - 1)) / rows
    for i, panel in enumerate(panels):
        row, col = divmod(i, cols)
        px = x + col * (cell_w + gap + (arrow_gap if rows == 1 else 0))
        py = y + row * (cell_h + gap)
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(px), Inches(py), Inches(cell_w), Inches(cell_h),
        )
        shape.fill.solid()
        shape.fill.fore_color.rgb = _PANEL_FILL
        shape.line.color.rgb = _PANEL_LINE
        tf = shape.text_frame
        tf.clear()
        _configure_text_frame(
            tf,
            margin_x=0.14,
            margin_y=0.08,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        heading = str(panel.get("heading", "")).strip()
        body = [str(v) for v in panel.get("body", []) if str(v).strip()]
        p = tf.paragraphs[0]
        _set_visible_paragraph(
            p,
            heading,
            font_size=18 if count <= 4 else 16,
            color=_ACCENT,
            bold=True,
            alignment=PP_ALIGN.CENTER,
        )
        p.space_after = Pt(6)
        body_chars = len(" ".join(body))
        for line in body:
            q = tf.add_paragraph()
            _set_visible_paragraph(
                q,
                line,
                font_size=_panel_text_size(count, body_chars),
            )
            q.space_after = Pt(3)
        if process and rows == 1 and col < cols - 1:
            ax = px + cell_w + 0.04
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(ax), Inches(py + cell_h / 2 - 0.12), Inches(0.24), Inches(0.24),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = _ACCENT
            arrow.line.color.rgb = _ACCENT


def _add_table(slide, headers: list[str], rows: list[list[str]], x: float, y: float, w: float, h: float) -> None:
    if not headers or not rows:
        return
    cols = len(headers)
    cleaned = []
    for row in rows[:10]:
        values = [str(v) for v in row[:cols]]
        values += [""] * (cols - len(values))
        cleaned.append(values)
    graphic = slide.shapes.add_table(len(cleaned) + 1, cols, Inches(x), Inches(y), Inches(w), Inches(h))
    table = graphic.table
    font_size = 17 if len(cleaned) <= 5 and cols <= 4 else 14
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.text_frame.clear()
        _configure_text_frame(
            cell.text_frame,
            margin_x=0.04,
            margin_y=0.02,
            vertical_anchor=MSO_ANCHOR.MIDDLE,
        )
        cell.fill.solid()
        cell.fill.fore_color.rgb = _PANEL_FILL
        p = cell.text_frame.paragraphs[0]
        _set_visible_paragraph(
            p,
            str(header),
            font_size=font_size,
            color=_ACCENT,
            bold=True,
            alignment=PP_ALIGN.CENTER,
        )
    for r, row in enumerate(cleaned, start=1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text_frame.clear()
            _configure_text_frame(
                cell.text_frame,
                margin_x=0.04,
                margin_y=0.02,
                vertical_anchor=MSO_ANCHOR.MIDDLE,
            )
            p = cell.text_frame.paragraphs[0]
            _set_visible_paragraph(p, value, font_size=font_size)


def _fit_picture(slide, image_path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(image_path) as image:
        px_w, px_h = image.size
    aspect = px_w / max(1, px_h)
    box_aspect = w / max(0.01, h)
    if aspect >= box_aspect:
        draw_w = w
        draw_h = w / aspect
    else:
        draw_h = h
        draw_w = h * aspect
    draw_x = x + (w - draw_w) / 2
    draw_y = y + (h - draw_h) / 2
    slide.shapes.add_picture(
        str(image_path), Inches(draw_x), Inches(draw_y), Inches(draw_w), Inches(draw_h)
    )


def _resolve_figure_paths(manifest_path: Path, item: dict) -> list[Path]:
    workspace = manifest_path.parent.parent
    paths: list[Path] = []
    for asset in item.get("figure_assets", []) or []:
        rel = str(asset.get("asset_path", "")).strip()
        if not rel:
            continue
        path = workspace / rel
        if path.is_file():
            paths.append(path)
    return paths


def _add_picture_grid(slide, paths: list[Path], x: float, y: float, w: float, h: float) -> None:
    if not paths:
        return
    count = len(paths)
    cols = 1 if count == 1 else 2 if count <= 4 else 3
    rows = math.ceil(count / cols)
    gap = 0.12
    cell_w = (w - gap * (cols - 1)) / cols
    cell_h = (h - gap * (rows - 1)) / rows
    for index, path in enumerate(paths):
        row, col = divmod(index, cols)
        _fit_picture(
            slide,
            path,
            x + col * (cell_w + gap),
            y + row * (cell_h + gap),
            cell_w,
            cell_h,
        )


def _add_notes(slide, item: dict) -> None:
    notes_tf = slide.notes_slide.notes_text_frame
    notes_tf.text = item.get("narration", "")
    guidance: list[str] = []
    if item.get("lecturer_notes"):
        guidance.append("LECTURER NOTES:\n" + "\n".join(f"- {x}" for x in item.get("lecturer_notes", [])))
    if item.get("visual_direction"):
        guidance.append("VISUAL DIRECTION:\n" + str(item.get("visual_direction")))
    if item.get("visual_type"):
        guidance.append("VISUAL TYPE: " + str(item.get("visual_type")))
    if item.get("source_block_ids"):
        guidance.append("SOURCE BLOCKS: " + ", ".join(item.get("source_block_ids", [])))
    if item.get("figure_ids"):
        guidance.append("TEXTBOOK FIGURES: " + ", ".join(item.get("figure_ids", [])))
    if guidance:
        p_notes = notes_tf.add_paragraph()
        p_notes.text = "\n\n".join(guidance)


def _build_standard_slide(prs: Presentation, item: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    _add_title(slide, item.get("title", ""))
    lines = [str(x) for x in item.get("onscreen", [])]
    visual_type = str(item.get("visual_type", "auto") or "auto")
    panels = item.get("visual_panels", []) or []
    headers = [str(x) for x in item.get("table_headers", []) or []]
    rows = item.get("table_rows", []) or []

    if visual_type == "table" and headers and rows:
        lead_h = _lead_region_height(lines, 11.4, preferred_font_size=24, maximum=1.65)
        if lead_h:
            _add_body(
                slide,
                item,
                0.95,
                _STANDARD_CONTENT_TOP,
                11.4,
                lead_h,
                has_figure=False,
                show_equation=False,
            )
        table_y = _STANDARD_CONTENT_TOP + lead_h + (_REGION_GAP if lead_h else 0.0)
        equation_h = 1.05 if item.get("equation_latex") else 0.0
        _add_table(
            slide,
            headers,
            rows,
            0.95,
            table_y,
            11.4,
            _SLIDE_BOTTOM - equation_h - table_y,
        )
        if item.get("equation_latex"):
            _add_equation_footer(slide, str(item["equation_latex"]), top=_SLIDE_BOTTOM - 0.95)
    elif visual_type == "equation_focus":
        _add_equation_focus(slide, item)
    elif visual_type in {"process", "comparison", "diagram"} and panels:
        equation_h = 1.05 if item.get("equation_latex") else 0.0
        content_bottom = _SLIDE_BOTTOM - equation_h
        lead_h = _lead_region_height(lines, 11.4, preferred_font_size=24, maximum=2.1)
        if lead_h:
            _add_body(
                slide,
                item,
                0.95,
                _STANDARD_CONTENT_TOP,
                11.4,
                lead_h,
                has_figure=False,
                show_equation=False,
            )
        panel_y = _STANDARD_CONTENT_TOP + lead_h + (_REGION_GAP if lead_h else 0.0)
        panel_h = content_bottom - panel_y
        _add_panel_grid(slide, panels, 0.95, panel_y, 11.4, panel_h, process=visual_type == "process")
        if item.get("equation_latex"):
            _add_equation_footer(slide, str(item["equation_latex"]), top=_SLIDE_BOTTOM - 0.95)
    else:
        _add_body(slide, item, 0.95, 1.55, 11.45, 5.35, has_figure=False)

    _add_notes(slide, item)
    return slide


def _build_textbook_title_slide(prs: Presentation, lesson_title: str, item: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(11.2), Inches(2.7))
    tf = box.text_frame
    tf.clear()
    _configure_text_frame(
        tf,
        margin_x=0.0,
        margin_y=0.0,
        vertical_anchor=MSO_ANCHOR.MIDDLE,
    )
    p = tf.paragraphs[0]
    n = len(lesson_title.strip())
    title_size = 36 if n <= 55 else 31 if n <= 85 else 26
    _set_visible_paragraph(
        p,
        lesson_title,
        font_size=title_size,
        bold=True,
        alignment=PP_ALIGN.CENTER,
    )
    _add_notes(slide, item)
    return slide


def _build_textbook_content_slide(prs: Presentation, manifest_path: Path, item: dict):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    figure_paths = _resolve_figure_paths(manifest_path, item)
    hint = str(item.get("figure_layout_hint", "auto"))
    lines = [str(x) for x in item.get("onscreen", [])]
    visual_type = str(item.get("visual_type", "auto") or "auto")
    panels = item.get("visual_panels", []) or []
    headers = [str(x) for x in item.get("table_headers", []) or []]
    rows = item.get("table_rows", []) or []

    if figure_paths:
        if len(figure_paths) == 1 and (hint == "image_large" or len(" ".join(lines)) <= 90):
            _add_picture_grid(slide, figure_paths, 1.0, 0.55, 11.3, 4.75)
            _add_body(slide, item, 1.1, 5.35, 11.1, 1.35, has_figure=True)
        elif hint == "image_left":
            _add_picture_grid(slide, figure_paths, 0.65, 0.75, 6.4, 5.95)
            _add_body(slide, item, 7.3, 0.95, 5.15, 5.6, has_figure=True)
        elif len(figure_paths) >= 2 or hint == "grid":
            _add_body(slide, item, 0.65, 0.9, 4.35, 5.7, has_figure=True)
            _add_picture_grid(slide, figure_paths, 5.2, 0.65, 7.45, 6.0)
        else:
            _add_body(slide, item, 0.65, 0.95, 5.15, 5.55, has_figure=True)
            _add_picture_grid(slide, figure_paths, 6.05, 0.75, 6.55, 5.95)
    elif visual_type == "table" and headers and rows:
        lead_h = _lead_region_height(lines, 11.4, preferred_font_size=24, maximum=1.65)
        if lead_h:
            _add_body(
                slide,
                item,
                0.95,
                _TEXTBOOK_CONTENT_TOP,
                11.4,
                lead_h,
                has_figure=False,
                show_equation=False,
            )
        table_y = _TEXTBOOK_CONTENT_TOP + lead_h + (_REGION_GAP if lead_h else 0.0)
        equation_h = 1.05 if item.get("equation_latex") else 0.0
        _add_table(
            slide,
            headers,
            rows,
            0.95,
            table_y,
            11.4,
            _SLIDE_BOTTOM - equation_h - table_y,
        )
        if item.get("equation_latex"):
            _add_equation_footer(slide, str(item["equation_latex"]), top=_SLIDE_BOTTOM - 0.95)
    elif visual_type == "equation_focus":
        _add_equation_focus(slide, item, content_top=_TEXTBOOK_CONTENT_TOP)
    elif visual_type in {"process", "comparison", "diagram"} and panels:
        equation_h = 1.05 if item.get("equation_latex") else 0.0
        lead_h = _lead_region_height(lines, 11.4, preferred_font_size=24, maximum=2.0)
        if lead_h:
            _add_body(
                slide,
                item,
                0.95,
                _TEXTBOOK_CONTENT_TOP,
                11.4,
                lead_h,
                has_figure=False,
                show_equation=False,
            )
        panel_y = _TEXTBOOK_CONTENT_TOP + lead_h + (_REGION_GAP if lead_h else 0.0)
        _add_panel_grid(
            slide,
            panels,
            0.95,
            panel_y,
            11.4,
            _SLIDE_BOTTOM - equation_h - panel_y,
            process=visual_type == "process",
        )
        if item.get("equation_latex"):
            _add_equation_footer(slide, str(item["equation_latex"]), top=_SLIDE_BOTTOM - 0.95)
    else:
        _add_body(slide, item, 0.95, 0.85, 11.45, 5.95, has_figure=False)

    _add_notes(slide, item)
    return slide


def validate_native_math_pptx(path: str | Path) -> int:
    """Reject a deck if mathematical notation escaped native Office Math."""
    native_zones = 0
    violations: list[str] = []
    with ZipFile(path) as archive:
        slide_names = sorted(
            name
            for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        for slide_name in slide_names:
            root = etree.fromstring(archive.read(slide_name))
            native_zones += len(root.findall(f".//{{{DRAWING_2010_NS}}}m"))
            for paragraph in root.findall(f".//{{{DRAWING_NS}}}p"):
                plain = "".join(
                    node.text or "" for node in paragraph.findall(f".//{{{DRAWING_NS}}}t")
                ).strip()
                if plain and (looks_like_math(plain) or re.search(r"\$[^$]+\$", plain)):
                    violations.append(f"{slide_name}: {plain[:120]}")
    if violations:
        raise RuntimeError(
            "Visible mathematical notation was not emitted as native Office Math: "
            + "; ".join(violations)
        )
    return native_zones


def build_pptx(manifest_path: str | Path, output_path: str | Path) -> Path:
    manifest_path = Path(manifest_path)
    m = yaml.safe_load(manifest_path.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    textbook = m.get("source_document_type") == "textbook_subchapter"

    for index, item in enumerate(m["slides"]):
        blocking = validate_slide_surface(item)
        if blocking:
            details = "; ".join(issue["message"] for issue in blocking)
            raise ValueError(
                f"Cannot build slide {item.get('id', index + 1)}: {details}"
            )
        if textbook and index == 0:
            _build_textbook_title_slide(prs, str(m.get("title", "")), item)
        elif textbook:
            _build_textbook_content_slide(prs, manifest_path, item)
        else:
            _build_standard_slide(prs, item)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output_path)
    validate_native_math_pptx(output_path)
    return output_path
