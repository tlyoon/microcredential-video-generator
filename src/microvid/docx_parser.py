from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from .models import Block

_DEFAULT_SECTION_RE = r"^(\d+(?:\.\d+)*)\.?\s+"
_DEFAULT_HEADING_STYLE_PATTERNS = [r"^Heading\s*(\d+)$"]


def _compiled_patterns(parser_config: dict | None, key: str, default: list[str]) -> list[re.Pattern]:
    values = (parser_config or {}).get(key, default)
    return [re.compile(str(v), re.IGNORECASE) for v in values]


def _section_from_heading(text: str, parser_config: dict | None = None) -> str | None:
    pattern = str((parser_config or {}).get("section_number_regex", _DEFAULT_SECTION_RE))
    m = re.match(pattern, text.strip())
    return m.group(1) if m and m.groups() else None


def _heading_level(paragraph, parser_config: dict | None = None) -> int | None:
    style_name = paragraph.style.name if paragraph.style else None
    if style_name:
        for pattern in _compiled_patterns(parser_config, "heading_style_patterns", _DEFAULT_HEADING_STYLE_PATTERNS):
            m = pattern.search(style_name)
            if m:
                if m.groups():
                    try:
                        return int(m.group(1))
                    except (TypeError, ValueError):
                        pass
                return 1

    # Custom Word styles may still expose an outline level even when their names are
    # not "Heading N". Word stores 0 for level 1, 1 for level 2, and so on.
    if (parser_config or {}).get("use_outline_level", True):
        nodes = paragraph._p.xpath("./w:pPr/w:outlineLvl")
        if nodes:
            try:
                return int(nodes[0].get(qn("w:val"))) + 1
            except (TypeError, ValueError):
                pass

    # Optional fallback for structurally numbered documents that do not use heading styles.
    if (parser_config or {}).get("infer_numbered_headings", False):
        text = paragraph.text.strip()
        sec = _section_from_heading(text, parser_config)
        if sec:
            return sec.count(".") + 1
    return None


def _matches_any(text: str, patterns: list[re.Pattern]) -> bool:
    return any(p.search(text) for p in patterns)


def _classify_paragraph(text: str, level: int | None, omml: list[str], parser_config: dict | None) -> str:
    if level is not None:
        return "heading"
    if not text and omml:
        return "equation"

    cfg = parser_config or {}
    key_patterns = _compiled_patterns(cfg, "key_idea_patterns", [r"^key\s+idea\b"])
    worked_patterns = _compiled_patterns(cfg, "worked_example_patterns", [r"^worked\s+example\b", r"^example\b"])
    contrast_patterns = _compiled_patterns(
        cfg,
        "contrast_label_patterns",
        [r"^poor:?$", r"^better:?$", r"^weak:?$", r"^incorrect\s+treatment:?$"],
    )

    if _matches_any(text, key_patterns):
        return "key_idea"
    if _matches_any(text, worked_patterns):
        return "worked_example_heading"
    if _matches_any(text, contrast_patterns):
        return "contrast_label"
    return "paragraph"


def _omml_nodes(paragraph) -> list:
    # Avoid returning both oMathPara and its child oMath, which would duplicate display equations.
    return paragraph._p.xpath(".//m:oMathPara | .//m:oMath[not(ancestor::m:oMathPara)]")


def _omml_xml(paragraph) -> list[str]:
    return [etree.tostring(node, encoding="unicode") for node in _omml_nodes(paragraph)]


def _omml_text(paragraph) -> list[str]:
    out: list[str] = []
    for node in _omml_nodes(paragraph):
        tokens = node.xpath(".//*[local-name()='t']/text()")
        text = "".join(tokens).replace("\u2005", " ").strip()
        if text:
            out.append(text)
    return out


def _iter_body_items(document: Document) -> Iterable[tuple[str, object]]:
    body = document.element.body
    p_map = {p._p: p for p in document.paragraphs}
    t_map = {t._tbl: t for t in document.tables}
    for child in body.iterchildren():
        if child.tag == qn("w:p") and child in p_map:
            yield "paragraph", p_map[child]
        elif child.tag == qn("w:tbl") and child in t_map:
            yield "table", t_map[child]


def extract_docx(path: str | Path, parser_config: dict | None = None) -> dict:
    """Extract DOCX structure without relying on pagination.

    The returned blocks carry both optional numeric section identifiers and a semantic
    heading path.  Downstream profiles can therefore select by section number for a
    tightly controlled document, or by heading text/path for documents whose numbering
    changes between revisions.
    """
    path = Path(path)
    doc = Document(path)
    blocks: list[Block] = []
    current_section: str | None = None
    sequence = 0
    heading_stack: dict[int, str] = {}

    for kind, item in _iter_body_items(doc):
        sequence += 1
        if kind == "paragraph":
            text = item.text.strip()
            style_name = item.style.name if item.style else None
            level = _heading_level(item, parser_config)
            math_text = _omml_text(item)
            omml = _omml_xml(item)

            if level is not None:
                # Update semantic breadcrumb independently of numeric section labels.
                heading_stack[level] = text
                for k in list(heading_stack):
                    if k > level:
                        del heading_stack[k]
                sec = _section_from_heading(text, parser_config)
                if sec:
                    current_section = sec
                elif level == 1:
                    # Do not leak a numeric section id from a preceding top-level section.
                    current_section = None

            if not text and not omml:
                continue

            block_kind = _classify_paragraph(text, level, omml, parser_config)
            if not text and math_text:
                text = " ; ".join(math_text)

            metadata = {
                "heading_path": [heading_stack[k] for k in sorted(heading_stack)],
            }
            if math_text:
                metadata["math_text"] = math_text

            blocks.append(
                Block(
                    id=f"b{len(blocks)+1:04d}",
                    kind=block_kind,
                    text=text,
                    section=current_section,
                    heading_level=level,
                    style=style_name,
                    source_index=sequence,
                    omml=omml,
                    metadata=metadata,
                )
            )
        else:
            rows = [[cell.text.strip() for cell in row.cells] for row in item.rows]
            text = "\n".join(" | ".join(row) for row in rows)
            blocks.append(
                Block(
                    id=f"b{len(blocks)+1:04d}",
                    kind="table",
                    text=text,
                    section=current_section,
                    source_index=sequence,
                    metadata={
                        "rows": rows,
                        "heading_path": [heading_stack[k] for k in sorted(heading_stack)],
                    },
                )
            )

    return {
        "schema_version": 2,
        "source": str(path),
        "block_count": len(blocks),
        "pagination_used": False,
        "blocks": [b.to_dict() for b in blocks],
    }


def write_extraction(
    path: str | Path,
    output: str | Path,
    parser_config: dict | None = None,
) -> dict:
    payload = extract_docx(path, parser_config=parser_config)
    output = Path(output)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return payload
