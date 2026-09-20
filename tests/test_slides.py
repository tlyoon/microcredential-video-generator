from zipfile import ZipFile

import pytest
import yaml
from lxml import etree
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE

from microvid.office_math import DRAWING_2010_NS, MATH_NS, looks_like_math, visible_math_to_latex
from microvid.slides import build_pptx, validate_native_math_pptx


def test_build_pptx(tmp_path):
    manifest = {
        "video_id": "V01",
        "title": "Measurement",
        "slides": [
            {
                "id": "V01S01",
                "slide_type": "hook",
                "title": "Measurement",
                "onscreen": ["Why uncertainty matters"],
                "source_sections": [],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "out.pptx"

    build_pptx(source, output)

    assert output.exists()
    assert len(Presentation(output).slides) == 1


def test_build_pptx_emits_native_office_math(tmp_path):
    manifest = {
        "video_id": "V01",
        "title": "Measurement",
        "slides": [
            {
                "id": "V01S01",
                "slide_type": "equation",
                "title": "Density",
                "onscreen": ["Mass: m = 56.20 ± 0.05 g"],
                "equation_latex": r"\rho = \frac{4m}{\pi D^2 h}",
                "source_sections": ["9.1"],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "math.pptx"

    build_pptx(source, output)

    with ZipFile(output) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml")
    root = etree.fromstring(slide_xml)
    assert len(root.findall(f".//{{{DRAWING_2010_NS}}}m")) == 2
    assert len(root.findall(f".//{{{MATH_NS}}}oMath")) == 2
    assert "Mass:\u00a0m\u00a0=" in "".join(root.itertext())
    assert b"Equation (LaTeX source)" not in slide_xml
    assert b"\\frac" not in slide_xml


def test_build_pptx_handles_overbar_as_native_office_math(tmp_path):
    manifest = {
        "video_id": "V03",
        "title": "Statistics",
        "slides": [
            {
                "id": "V03S01",
                "slide_type": "equation",
                "title": "Mean",
                "onscreen": [],
                "equation_latex": r"\bar{x} = \frac{1}{N}\sum_{i=1}^{N}x_i",
                "source_sections": ["4.1"],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "bar.pptx"

    build_pptx(source, output)

    with ZipFile(output) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml")
    root = etree.fromstring(slide_xml)
    assert len(root.findall(f".//{{{DRAWING_2010_NS}}}m")) == 1
    assert len(root.findall(f".//{{{MATH_NS}}}bar")) == 1


def test_visible_math_normalizes_ascii_and_unicode_notation():
    latex = visible_math_to_latex("Mean: D_bar = sqrt(sum(x_i)^2); T²; t₁₀ approx 2")

    assert r"\bar{D}" in latex
    assert r"\sqrt{\sum" in latex
    assert r"\sum" in latex
    assert "T^{2}" in latex
    assert "t_{10}" in latex
    assert r"\approx" in latex
    assert r"\sqrt{a^2+b^2}" in visible_math_to_latex("sqrt[a^2+b^2]")


def test_standard_pptx_renders_visual_panels_without_provenance_footer(tmp_path):
    manifest = {
        "video_id": "V01",
        "title": "Compare two ideas",
        "slides": [
            {
                "id": "V01S01",
                "slide_type": "introduction",
                "title": "Compare two ideas",
                "onscreen": ["One question, two possible interpretations"],
                "visual_type": "comparison",
                "visual_panels": [
                    {"heading": "Case A", "body": ["First interpretation"]},
                    {"heading": "Case B", "body": ["Second interpretation"]},
                ],
                "table_headers": [],
                "table_rows": [],
                "source_sections": ["1.1"],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "visual.pptx"

    build_pptx(source, output)

    prs = Presentation(output)
    text = " ".join(shape.text for shape in prs.slides[0].shapes if hasattr(shape, "text"))
    assert "Case A" in text
    assert "Case B" in text
    assert "source sections" not in text
    assert "V01S01" not in text


def test_standard_pptx_renders_source_grounded_table(tmp_path):
    manifest = {
        "video_id": "V02",
        "title": "Compare measurements",
        "slides": [
            {
                "id": "V02S01",
                "slide_type": "concept",
                "title": "Compare measurements",
                "onscreen": ["Alignment makes the difference easy to inspect."],
                "visual_type": "table",
                "visual_panels": [],
                "table_headers": ["Quantity", "Meaning"],
                "table_rows": [["Precision", "repeatability"], ["Accuracy", "closeness to reference"]],
                "source_sections": ["2.1"],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "table.pptx"

    build_pptx(source, output)

    prs = Presentation(output)
    table_shapes = [shape for shape in prs.slides[0].shapes if getattr(shape, "has_table", False)]
    assert len(table_shapes) == 1
    table = table_shapes[0].table
    assert table.cell(0, 0).text == "Quantity"
    assert table.cell(2, 1).text == "closeness to reference"


def test_textbook_long_title_uses_adaptive_font_size(tmp_path):
    title = "The Nonisolated System: Energy Transfers and the Master Equation for a Carefully Defined Physical System"
    manifest = {
        "video_id": "V01",
        "title": title,
        "source_document_type": "textbook_subchapter",
        "slides": [
            {
                "id": "V01S01",
                "slide_type": "title",
                "title": title,
                "onscreen": [],
                "visual_type": "text",
                "visual_panels": [],
                "table_headers": [],
                "table_rows": [],
                "source_sections": ["8.1"],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "long-title.pptx"

    build_pptx(source, output)

    prs = Presentation(output)
    title_shape = next(shape for shape in prs.slides[0].shapes if getattr(shape, "has_text_frame", False) and title in shape.text)
    run = title_shape.text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt <= 26
    assert title_shape.left >= 0
    assert title_shape.left + title_shape.width <= prs.slide_width


def test_visual_slide_allocates_non_overlapping_autofit_regions(tmp_path):
    manifest = {
        "video_id": "V03",
        "title": "Safe layout",
        "slides": [
            {
                "id": "V03S01",
                "slide_type": "concept",
                "title": "A dense lead must not sit behind panels",
                "onscreen": [
                    "First explanatory line may wrap across the available width.",
                    "Second line still needs a separate lead-text region.",
                ],
                "visual_type": "comparison",
                "visual_panels": [
                    {"heading": "Case A", "body": ["First interpretation"]},
                    {"heading": "Case B", "body": ["Second interpretation"]},
                ],
            }
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest), encoding="utf-8")
    output = tmp_path / "safe-layout.pptx"

    build_pptx(source, output)

    slide = Presentation(output).slides[0]
    lead = next(shape for shape in slide.shapes if "First explanatory line" in getattr(shape, "text", ""))
    panels = [shape for shape in slide.shapes if "Case " in getattr(shape, "text", "")]
    assert lead.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    assert panels
    assert lead.top + lead.height <= min(panel.top for panel in panels)
    assert all(panel.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE for panel in panels)


def test_panels_and_tables_emit_detected_notation_as_native_office_math(tmp_path):
    manifest = {
        "video_id": "V06",
        "title": "Native mathematics",
        "slides": [
            {
                "id": "V06S01",
                "slide_type": "concept",
                "title": "Panel mathematics",
                "onscreen": ["Compare the two source-supported relations."],
                "visual_type": "comparison",
                "visual_panels": [
                    {"heading": "Mass", "body": ["m = 2.0 kg"]},
                    {"heading": "Speed", "body": ["v² = 9 m² s⁻²"]},
                ],
            },
            {
                "id": "V06S02",
                "slide_type": "concept",
                "title": "Table mathematics",
                "onscreen": ["Keep measured values editable."],
                "visual_type": "table",
                "table_headers": ["Quantity", "Value"],
                "table_rows": [["Acceleration", "g = 9.81 m s⁻²"]],
            },
        ],
    }
    source = tmp_path / "manifest.yaml"
    source.write_text(yaml.safe_dump(manifest, allow_unicode=True), encoding="utf-8")
    output = tmp_path / "all-native-math.pptx"

    build_pptx(source, output)

    with ZipFile(output) as archive:
        roots = [
            etree.fromstring(archive.read("ppt/slides/slide1.xml")),
            etree.fromstring(archive.read("ppt/slides/slide2.xml")),
        ]
    assert len(roots[0].findall(f".//{{{DRAWING_2010_NS}}}m")) == 2
    assert len(roots[1].findall(f".//{{{DRAWING_2010_NS}}}m")) == 1
    assert looks_like_math("x₂ ≤ 3")
    assert looks_like_math("ρ = m/V")
    assert looks_like_math("12.4 ± 0.1 cm")


def test_native_math_invariant_rejects_plain_text_equations(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 4_000_000, 1_000_000)
    box.text_frame.paragraphs[0].text = "rho = m / V"
    output = tmp_path / "plain-math.pptx"
    presentation.save(output)

    with pytest.raises(RuntimeError, match="not emitted as native Office Math"):
        validate_native_math_pptx(output)
