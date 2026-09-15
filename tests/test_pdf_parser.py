from pathlib import Path

import fitz
import pytest
from docx import Document

from microvid.pdf_parser import extract_pdf
from microvid.source_parser import extract_source


def _write_pdf(path: Path) -> None:
    doc = fitz.open()
    for page_no in (1, 2, 3):
        page = doc.new_page(width=612, height=792)
        page.insert_text((72, 36), "COURSE HEADER", fontsize=8)
        page.insert_text((280, 750), f"Page {page_no}", fontsize=8)
        if page_no == 1:
            page.insert_text((72, 90), "1. Measurement", fontsize=16, fontname="hebo")
            page.insert_text((72, 120), "A measured value needs a unit and an uncertainty.", fontsize=11)
            page.insert_text((72, 155), "1.1 Resolution", fontsize=14, fontname="hebo")
            page.insert_text((72, 185), "Instrument resolution is not the whole uncertainty.", fontsize=11)
        elif page_no == 2:
            page.insert_text((72, 90), "2. Propagation", fontsize=16, fontname="hebo")
            page.insert_text((72, 120), "A calculated result inherits uncertainty from its inputs.", fontsize=11)
            page.insert_text((72, 155), "Worked example: speed", fontsize=12, fontname="hebo")
            page.insert_text((72, 185), "v = d / t", fontsize=11, fontname="symb")
        else:
            page.insert_text((72, 90), "3. Reporting", fontsize=16, fontname="hebo")
            page.insert_text((72, 120), "Report the final value with uncertainty and units.", fontsize=11)
    doc.save(path)
    doc.close()


def test_extract_pdf_semantic_blocks_and_provenance(tmp_path):
    path = tmp_path / "sample.pdf"
    _write_pdf(path)

    payload = extract_pdf(path)

    assert payload["source_format"] == "pdf"
    assert payload["page_count"] == 3
    assert payload["pagination_used"] is False
    assert not any("COURSE HEADER" in block["text"] for block in payload["blocks"])
    assert not any(block["text"].startswith("Page ") for block in payload["blocks"])
    assert any(
        block["kind"] == "heading" and block["section"] == "1.1" and block["heading_level"] == 2
        for block in payload["blocks"]
    )
    worked = next(block for block in payload["blocks"] if block["text"] == "Worked example: speed")
    assert worked["kind"] == "worked_example_heading"
    assert worked["metadata"]["page_number"] == 2
    assert worked["metadata"]["heading_path"][-1] == "2. Propagation"


def test_source_dispatcher_keeps_docx_support(tmp_path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_heading("1. Measurement", level=1)
    doc.add_paragraph("A measured value needs a unit.")
    doc.save(path)

    payload = extract_source(path)
    assert payload["source_format"] == "docx"
    assert any(block["section"] == "1" for block in payload["blocks"])


def test_source_dispatcher_accepts_pdf_and_rejects_unknown(tmp_path):
    pdf = tmp_path / "sample.pdf"
    _write_pdf(pdf)
    assert extract_source(pdf)["source_format"] == "pdf"

    bad = tmp_path / "sample.txt"
    bad.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError, match="Unsupported source format"):
        extract_source(bad)


def test_image_only_pdf_fails_visibly(tmp_path):
    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page(width=612, height=792)
    doc.save(path)
    doc.close()

    with pytest.raises(ValueError, match="too little extractable text"):
        extract_pdf(path)


def test_write_extraction_serializes_common_schema(tmp_path):
    from microvid.source_parser import write_extraction

    source = tmp_path / "sample.pdf"
    output = tmp_path / "workspace" / "extracted" / "document_structure.json"
    _write_pdf(source)
    payload = write_extraction(source, output)

    import json

    saved = json.loads(output.read_text(encoding="utf-8"))
    assert saved["block_count"] == payload["block_count"]
    assert saved["source_format"] == "pdf"
    assert all("heading_path" in block["metadata"] for block in saved["blocks"])


def test_legacy_cli_extraction_entry_point_accepts_pdf(tmp_path):
    from microvid.docx_parser import write_extraction as cli_write_extraction

    source = tmp_path / "sample.pdf"
    output = tmp_path / "workspace" / "extracted" / "document_structure.json"
    _write_pdf(source)
    payload = cli_write_extraction(source, output)

    assert payload["source_format"] == "pdf"
    assert output.is_file()
