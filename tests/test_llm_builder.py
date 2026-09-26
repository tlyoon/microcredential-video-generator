from docx import Document

from microvid.docx_parser import extract_docx
from microvid.llm_manifest_builder import build_lesson_manifest_with_llm
from microvid.qa import validate_manifest


class FakeProvider:
    provider_name = "gemini"
    model = "fake-gemini"

    def __init__(self):
        self.calls = 0
        self.prompts = []

    def generate_json(self, prompt, schema):
        self.calls += 1
        self.prompts.append(prompt)
        if "DEDICATED NARRATION POLISH" in prompt:
            return {
                "slides": [
                    {
                        "slide_id": "V01S01",
                        "narration": "Suppose someone writes down twelve point four and calls it a measurement. What, exactly, have they measured? Until we attach a unit, the number has no physical scale.",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S02",
                        "narration": "Now attach the unit. The value twelve point four becomes meaningful as a length only when we report it in millimetres. The unit is part of the measurement, not an optional label added afterward.",
                        "tts_text": "Now attach the unit. The value twelve point four becomes meaningful as a length only when we report it in millimetres. The unit is part of the measurement, not an optional label added afterward.",
                    },
                    {
                        "slide_id": "V01S03",
                        "narration": "The important distinction is between a bare number and a physical quantity. The same number could describe a length, a time, or something else. The unit tells us which physical scale the value belongs to.",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S04",
                        "narration": "Pause here. If someone reports only twelve point four, what essential information is still missing before you can interpret the measurement?",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S05",
                        "narration": "So the habit to keep is simple. A measured number becomes a physical result only when its unit travels with it. Whenever you record a value, ask whether another person could tell what physical quantity that number represents.",
                        "tts_text": None,
                    },
                ]
            }
        return {
            "lesson_title": "Measurement Basics",
            "learning_outcomes": ["Explain why a measurement needs a unit."],
            "slides": [
                {
                    "slide_type": "introduction",
                    "title": "Measurement Basics",
                    "onscreen": ["12.4 — what does this number mean?"],
                    "narration": "A laboratory result is more than a number. We need enough context to know what the number means.",
                    "lecturer_notes": ["Open with the ambiguity of a bare number."],
                    "visual_direction": "Contrast a bare 12.4 with a question mark and a labelled quantity.",
                    "visual_type": "comparison",
                    "visual_panels": [
                        {"heading": "Bare number", "body": ["12.4", "Physical meaning unknown"]},
                        {"heading": "Measurement", "body": ["12.4 mm", "Value + unit"]},
                    ],
                    "table_headers": [],
                    "table_rows": [],
                    "equation_latex": None,
                    "source_block_ids": ["b0001", "b0002"],
                    "estimated_seconds": 35,
                },
                {
                    "slide_type": "concept",
                    "title": "Value plus unit",
                    "onscreen": ["12.4 mm"],
                    "narration": "The source states that a measured value needs a unit, so the unit is part of the scientific meaning of the result.",
                    "lecturer_notes": ["Stress that units are not optional labels."],
                    "visual_direction": "Highlight the value and unit as two inseparable parts.",
                    "visual_type": "diagram",
                    "visual_panels": [
                        {"heading": "Value", "body": ["12.4"]},
                        {"heading": "Unit", "body": ["millimetres"]},
                    ],
                    "table_headers": [],
                    "table_rows": [],
                    "equation_latex": None,
                    "source_block_ids": ["b0001", "b0002"],
                    "estimated_seconds": 50,
                },
                {
                    "slide_type": "interpretation",
                    "title": "A number becomes a physical quantity",
                    "onscreen": ["Unit gives the numerical value a physical scale."],
                    "narration": "The unit turns an otherwise ambiguous number into a physical quantity that another person can interpret.",
                    "lecturer_notes": ["Connect notation to communicable physical meaning."],
                    "visual_direction": "Show the bare number moving into a labelled physical quantity.",
                    "visual_type": "process",
                    "visual_panels": [
                        {"heading": "Number", "body": ["12.4"]},
                        {"heading": "Attach unit", "body": ["millimetres"]},
                        {"heading": "Physical result", "body": ["12.4 mm"]},
                    ],
                    "table_headers": [],
                    "table_rows": [],
                    "equation_latex": None,
                    "source_block_ids": ["b0001", "b0002"],
                    "estimated_seconds": 45,
                },
                {
                    "slide_type": "check",
                    "title": "Check your understanding",
                    "onscreen": ["What is missing from 12.4?"],
                    "narration": "Pause and identify what information is needed before this can be treated as a measurement.",
                    "lecturer_notes": ["Pause before discussion."],
                    "visual_direction": "Keep the question large and uncluttered.",
                    "visual_type": "text",
                    "visual_panels": [],
                    "table_headers": [],
                    "table_rows": [],
                    "equation_latex": None,
                    "source_block_ids": [],
                    "estimated_seconds": 30,
                },
                {
                    "slide_type": "conclusion",
                    "title": "A measurement needs physical meaning",
                    "onscreen": ["Measured result = value + unit"],
                    "narration": "A measurement becomes interpretable only when the value and its unit are reported together.",
                    "lecturer_notes": ["Resolve the opening question in one sentence."],
                    "visual_direction": "Return to 12.4 mm with value and unit visually joined.",
                    "visual_type": "diagram",
                    "visual_panels": [
                        {"heading": "Keep together", "body": ["numerical value", "unit"]}
                    ],
                    "table_headers": [],
                    "table_rows": [],
                    "equation_latex": None,
                    "source_block_ids": ["b0001", "b0002"],
                    "estimated_seconds": 30,
                },
            ],
            "editorial_flags": [],
        }


def test_llm_builder_adds_dedicated_narration_polish_and_preserves_provenance(tmp_path):
    p = tmp_path / "sample.docx"
    d = Document()
    d.add_heading("1. Measurement", level=1)
    d.add_paragraph("A measured value needs a unit.")
    d.save(p)

    extraction = extract_docx(p)
    course = {
        "title": "Test",
        "audience": "Freshmen",
        "narration_wpm": 130,
        "max_slides": 7,
        "llm": {"narration_polish_pass": True},
    }
    lesson = {
        "id": "V01",
        "title": "Measurement",
        "focus": "Understand measurement reporting.",
        "target_minutes": 3,
        "core_selectors": [{"heading_contains": "Measurement"}],
    }

    provider = FakeProvider()
    manifest = build_lesson_manifest_with_llm(
        extraction, course, lesson, provider, review_pass=True
    )

    assert provider.calls == 3
    assert any("DEDICATED NARRATION POLISH" in p for p in provider.prompts)
    assert any("Feynman-inspired" in p for p in provider.prompts)
    assert manifest["generation"]["mode"] == "llm"
    assert manifest["generation"]["passes"] == 3
    assert manifest["generation"]["narration_polish"]["enabled"] is True
    assert manifest["editorial_status"] == "automated_ready"
    assert manifest["slides"][0]["slide_type"] == "introduction"
    assert manifest["slides"][-1]["slide_type"] == "conclusion"
    assert manifest["slides"][1]["source_block_ids"] == ["b0001", "b0002"]
    assert "unit is part of the measurement" in manifest["slides"][1]["narration"]
    assert manifest["slides"][1]["tts_text"].startswith("Now attach the unit")
    assert manifest["slides"][1]["narration_word_count"] > 0
    assert not any(i["severity"] == "error" for i in validate_manifest(manifest))
