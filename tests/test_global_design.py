from docx import Document
import yaml

from microvid.docx_parser import extract_docx
from microvid.global_planner import build_global_course_plan, extraction_signature
from microvid.llm_manifest_builder import build_all_manifests_with_llm


class GlobalFakeProvider:
    provider_name = "gemini"
    model = "fake-global-gemini"

    def __init__(self):
        self.prompts = []

    def generate_json(self, prompt, schema):
        self.prompts.append(prompt)
        if "DEDICATED NARRATION POLISH PASS" in prompt:
            return {
                "slides": [
                    {
                        "slide_id": "V01S01",
                        "narration": "Start with two different questions. One asks how much energy can be stored, while the other asks how effectively supplied input becomes useful output. Keeping those questions separate will organize the whole lesson.",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S02",
                        "narration": "Capacity answers the storage question. It describes how much can be stored under the stated conditions, so those conditions are part of what gives the capacity its meaning.",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S03",
                        "narration": "Efficiency asks something different. Instead of how much can be stored, it compares the useful output with the supplied input. That distinction prevents us from treating capacity and efficiency as interchangeable ideas.",
                        "tts_text": None,
                    },
                    {
                        "slide_id": "V01S04",
                        "narration": "Pause here and decide which concept answers this question: useful output compared with supplied input. Focus on the type of comparison being made.",
                        "tts_text": None,
                    },
                ]
            }
        if "WHOLE-COURSE CONSISTENCY REVIEW" in prompt:
            return {
                "status": "ready",
                "summary": "The generated lesson follows the global plan coherently.",
                "findings": [],
                "lesson_revision_instructions": [],
            }
        if "GLOBAL COURSE PLANNING" in prompt or "GLOBAL COURSE PLAN REVIEW" in prompt:
            return {
                "course_summary": "Energy storage and efficiency form one short conceptual sequence.",
                "pedagogical_strategy": "Establish storage and capacity before comparing useful output with supplied input.",
                "concept_map": [
                    {
                        "name": "Energy storage",
                        "description": "Energy can be stored in physical systems.",
                        "source_block_ids": ["b0001", "b0002"],
                        "prerequisite_concepts": [],
                    },
                    {
                        "name": "Efficiency",
                        "description": "Efficiency compares useful output with supplied input.",
                        "source_block_ids": ["b0003", "b0004"],
                        "prerequisite_concepts": ["Energy storage"],
                    },
                ],
                "videos": [
                    {
                        "id": "V01",
                        "title": "From Storage to Efficiency",
                        "focus": "Connect storage capacity with efficiency as a coherent introductory lesson.",
                        "target_minutes": 3,
                        "max_slides": 6,
                        "learning_outcomes": ["Explain the two source concepts and their relationship."],
                        "check_question": "What does efficiency compare?",
                        "takeaways": ["Storage and efficiency answer different questions."],
                        "core_block_ids": ["b0001", "b0002", "b0003", "b0004"],
                        "reference_block_ids": [],
                        "prerequisite_video_ids": [],
                        "already_taught": [],
                        "forward_links": [],
                    }
                ],
                "coverage_notes": [],
                "editorial_flags": [],
            }
        return {
            "lesson_title": "From Storage to Efficiency",
            "learning_outcomes": ["Explain the two source concepts and their relationship."],
            "slides": [
                {
                    "slide_type": "hook",
                    "title": "Two questions about energy",
                    "onscreen": ["How much can be stored?", "How much useful output do we obtain?"],
                    "narration": "This lesson connects two questions that belong in one conceptual sequence.",
                    "lecturer_notes": ["Use the two questions as the organizing contrast."],
                    "visual_direction": "Show two labeled question cards.",
                    "equation_latex": None,
                    "source_block_ids": [],
                    "estimated_seconds": 30,
                },
                {
                    "slide_type": "concept",
                    "title": "Storage and capacity",
                    "onscreen": ["Capacity: how much can be stored under stated conditions."],
                    "narration": "The source describes capacity as how much can be stored under the stated conditions.",
                    "lecturer_notes": ["Keep the stated conditions visible."],
                    "visual_direction": "Highlight capacity beside a storage icon.",
                    "equation_latex": None,
                    "source_block_ids": ["b0001", "b0002"],
                    "estimated_seconds": 45,
                },
                {
                    "slide_type": "concept",
                    "title": "Efficiency asks a different question",
                    "onscreen": ["Efficiency compares useful output with supplied input."],
                    "narration": "Efficiency is about the comparison between useful output and supplied input.",
                    "lecturer_notes": ["Contrast this with capacity rather than treating them as synonyms."],
                    "visual_direction": "Show input and useful output with an arrow.",
                    "equation_latex": None,
                    "source_block_ids": ["b0003", "b0004"],
                    "estimated_seconds": 45,
                },
                {
                    "slide_type": "check",
                    "title": "Check the distinction",
                    "onscreen": ["Which idea answers: useful output compared with supplied input?"],
                    "narration": "Pause and choose which of the two concepts answers this question.",
                    "lecturer_notes": ["Pause briefly."],
                    "visual_direction": "Display only the question.",
                    "equation_latex": None,
                    "source_block_ids": [],
                    "estimated_seconds": 30,
                },
            ],
            "editorial_flags": [],
        }


def _make_source(path):
    d = Document()
    d.add_heading("1. Energy Storage", level=1)
    d.add_paragraph("Capacity describes how much energy can be stored under stated conditions.")
    d.add_heading("2. Efficiency", level=1)
    d.add_paragraph("Efficiency compares useful output with supplied input at the end of the process.")
    d.save(path)


def _profile():
    return {
        "course": {
            "id": "energy",
            "title": "Energy Course",
            "audience": "First-year students",
            "design_principle": "Video teaches reasoning; source carries detail.",
            "narration_wpm": 130,
            "target_video_minutes": 3,
            "max_slides": 6,
            "llm": {
                "narration_polish_pass": True,
                "global_design": {"max_source_characters": 100000, "max_videos": 10},
            },
        },
        "videos": [],
    }


def test_global_planner_reads_the_complete_document_before_segmentation(tmp_path):
    source = tmp_path / "source.docx"
    _make_source(source)
    extraction = extract_docx(source)
    provider = GlobalFakeProvider()

    plan = build_global_course_plan(extraction, _profile(), provider, review_pass=True)

    assert len(provider.prompts) == 2
    assert "Capacity describes how much energy can be stored" in provider.prompts[0]
    assert "Efficiency compares useful output with supplied input" in provider.prompts[0]
    assert plan["design_mode"] == "global_llm"
    assert plan["source_signature"] == extraction_signature(extraction)
    assert plan["videos"][0]["core_block_ids"] == ["b0001", "b0002", "b0003", "b0004"]


def test_per_lesson_generation_receives_global_map_polish_and_final_course_review(tmp_path):
    source = tmp_path / "source.docx"
    _make_source(source)
    extraction = extract_docx(source)
    provider = GlobalFakeProvider()
    profile = _profile()
    plan = build_global_course_plan(extraction, profile, provider, review_pass=True)

    paths = build_all_manifests_with_llm(
        extraction,
        profile,
        tmp_path / "manifests",
        provider,
        review_pass=True,
        global_plan=plan,
        course_consistency_review=True,
    )

    assert len(paths) == 1
    lesson_prompts = [p for p in provider.prompts if "CURRENT LESSON INPUT" in p]
    assert lesson_prompts
    assert "global_course_context" in lesson_prompts[0]
    assert "Energy storage and efficiency form one short conceptual sequence" in lesson_prompts[0]
    assert any("DEDICATED NARRATION POLISH PASS" in p for p in provider.prompts)
    assert any("WHOLE-COURSE CONSISTENCY REVIEW" in p for p in provider.prompts)

    manifest = yaml.safe_load(paths[0].read_text())
    assert manifest["generation"]["narration_polish"]["enabled"] is True
    assert manifest["slides"][0]["narration"].startswith("Start with two different questions")

    index = yaml.safe_load((tmp_path / "manifests" / "course.yaml").read_text())
    assert index["design_mode"] == "global_llm"
    assert index["global_consistency_status"] == "ready"
    assert index["narration_polish_enabled"] is True
