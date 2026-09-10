import yaml
from pptx import Presentation
from microvid.slides import build_pptx

def test_pptx_contains_narration_in_speaker_notes(tmp_path):
    manifest={"video_id":"V01","title":"Test","slides":[{"id":"V01S01","slide_type":"concept","title":"A concept","onscreen":["Short text"],"narration":"This is the narration script for this slide.","lecturer_notes":["Emphasize the distinction."],"visual_direction":"Highlight the key term.","source_sections":["1"],"source_block_ids":["b0001"]}]}
    m=tmp_path/"m.yaml"; m.write_text(yaml.safe_dump(manifest),encoding="utf-8"); out=tmp_path/"out.pptx"; build_pptx(m,out)
    text=Presentation(out).slides[0].notes_slide.notes_text_frame.text
    assert "This is the narration script" in text; assert "LECTURER NOTES" in text; assert "VISUAL DIRECTION" in text
