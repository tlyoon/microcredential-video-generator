from docx import Document
from microvid.docx_parser import extract_docx


def test_extract_headings_and_table(tmp_path):
    p = tmp_path / "sample.docx"
    d = Document()
    d.add_heading("1. Measurement", level=1)
    d.add_paragraph("A measured value needs a unit.")
    d.add_heading("1.1 Resolution", level=2)
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Instrument"
    table.cell(0, 1).text = "Resolution"
    table.cell(1, 0).text = "Ruler"
    table.cell(1, 1).text = "1 mm"
    d.save(p)

    payload = extract_docx(p)
    assert payload["block_count"] >= 4
    assert any(b["kind"] == "heading" and b["section"] == "1" for b in payload["blocks"])
    assert any(b["kind"] == "table" and "Ruler" in b["text"] for b in payload["blocks"])
