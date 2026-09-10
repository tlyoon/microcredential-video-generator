from docx import Document

from microvid.auto_profile import scaffold_profile
from microvid.docx_parser import extract_docx
from microvid.manifest_builder import build_lesson_manifest
from microvid.qa import validate_manifest
from microvid.segmenter import selector_matches


def _make_doc(path, prefix="3"):
    d = Document()
    d.add_heading(f"{prefix}. Energy Storage", level=1)
    d.add_paragraph(
        "Energy may be stored in several physical forms. The useful model depends on the system."
    )
    d.add_heading(f"{prefix}.1 Capacity", level=2)
    d.add_paragraph("Capacity describes how much can be stored under the stated conditions.")
    d.add_heading(f"{int(prefix)+1}. Efficiency", level=1)
    d.add_paragraph("Efficiency compares useful output with supplied input.")
    d.save(path)


def test_semantic_heading_selector_survives_renumbering(tmp_path):
    p1 = tmp_path / "a.docx"
    p2 = tmp_path / "b.docx"
    _make_doc(p1, "3")
    _make_doc(p2, "8")

    a = extract_docx(p1)
    b = extract_docx(p2)
    selector = {"heading_contains": "Energy Storage"}
    a_matches = [x for x in a["blocks"] if selector_matches(x, selector)]
    b_matches = [x for x in b["blocks"] if selector_matches(x, selector)]
    assert a_matches and b_matches
    assert any(x["section"] == "3" for x in a_matches)
    assert any(x["section"] == "8" for x in b_matches)
    assert a["pagination_used"] is False and b["pagination_used"] is False


def test_scaffold_profile_for_unrelated_topic_is_global_first(tmp_path):
    p = tmp_path / "topic.docx"
    _make_doc(p, "1")
    extraction = extract_docx(p)
    profile = scaffold_profile(
        extraction, "energy_course", "Energy Course", target_video_minutes=5
    )

    # The scaffold no longer guesses video boundaries from Heading-1 sections.
    assert profile["videos"] == []
    assert profile["course"]["llm"]["global_design"]["enabled"] is True
    assert profile["course"]["content_selection"]["priority_terms"] == []
    assert profile["course"]["profile_status"].startswith("global_design")


def test_profile_drift_is_error():
    extraction = {
        "blocks": [
            {
                "id": "b1",
                "kind": "heading",
                "text": "1. New Topic",
                "section": "1",
                "heading_level": 1,
                "source_index": 1,
                "metadata": {"heading_path": ["1. New Topic"]},
            }
        ]
    }
    course = {"narration_wpm": 130, "max_slides": 7}
    lesson = {
        "id": "V01",
        "title": "Old Topic",
        "focus": "Old focus",
        "core_selectors": [{"heading_contains": "Missing Heading"}],
        "target_minutes": 5,
    }
    manifest = build_lesson_manifest(extraction, course, lesson)
    issues = validate_manifest(manifest)
    assert any(
        i["severity"] == "error" and "zero authoritative source blocks" in i["message"]
        for i in issues
    )
